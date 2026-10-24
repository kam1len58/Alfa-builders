"""Оценка важности новостей для компаний по бизнес-матрице.

Запуск: python analytics/score_news.py
Берёт data/golden_news.json, data/demo_profiles.json, data/business_matrix.json.
Результат: data/feed_preview.xlsx (баллы и объяснения), data/feed_preview.json (то же для бэкенда)
и обновлённое поле matched_rules в golden_news.json.

Балл новости для компании: score = 1 - prod(1 - w_i) по сработавшим правилам.
score >= 0.6 — «В фокусе», 0.4–0.6 — «Все события», ниже — не показываем.
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATA = Path(__file__).resolve().parent.parent / "data"
FOCUS, ALL = 0.6, 0.4


def load(name):
    return json.load(open(DATA / name, encoding="utf-8"))


def profile_value(profile, field, field_map):
    path = field_map.get(field, field)
    value = profile
    for key in path.split("."):
        value = value.get(key) if isinstance(value, dict) else None
    return value


def condition_ok(profile, cond, field_map):
    v = profile_value(profile, cond["field"], field_map)
    if "okved_prefix" in cond:
        return bool(v) and any(str(v).startswith(p) for p in cond["okved_prefix"])
    if "eq" in cond:
        return v == cond["eq"]
    if "gte" in cond:
        return v is not None and v >= cond["gte"]
    if "in" in cond:
        return v is not None and any(x.lower() in str(v).lower() for x in cond["in"])
    return False  # совпадение отрасли контрагента с отраслью новости определяет ML


def stem(word):
    word = word.lower()
    return word if len(word) <= 4 else word[: max(4, len(word) - 2)]


def rule_triggered(rule, text):
    text = text.lower()
    for phrase in rule["triggers"]:
        pattern = r"\b" + r"\w*\s+".join(re.escape(stem(w)) for w in phrase.split()) + r"\w*"
        if re.search(pattern, text):
            return True
    return False


def bucket(score):
    return "В фокусе" if score >= FOCUS else "Все события" if score >= ALL else "Не показывать"


def main():
    matrix = load("business_matrix.json")
    golden = load("golden_news.json")
    companies = load("demo_profiles.json")
    field_map = matrix["field_map_csv_to_profile"]

    feed = []
    for news in golden["news"]:
        news["matched_rules"] = [r["id"] for r in matrix["rules"] if rule_triggered(r, news["what_happened"])]
        for company in companies:
            fired = [r for r in matrix["rules"] if r["id"] in news["matched_rules"]
                     and all(condition_ok(company, c, field_map) for c in r["condition"])]
            rest = 1.0
            for r in fired:
                rest *= 1 - r["weight"]
            score = round(1 - rest, 3)
            feed.append({
                "news_id": news["id"], "company_id": company["synthetic_u"], "scenario": company["demo_scenario"],
                "score": score, "bucket": bucket(score),
                "rules": [r["id"] for r in fired],
                "why": [r["prompt_hint"] for r in fired],
                "products": sorted({p for r in fired for p in r["bank_products"]}),
            })

    json.dump(golden, open(DATA / "golden_news.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    json.dump(feed, open(DATA / "feed_preview.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    wb = Workbook()
    ws = wb.active
    ws.title = "Баллы"
    ws.append(["Новость"] + [f"{c['synthetic_u']}\n{c['demo_scenario']}" for c in companies])
    for news in golden["news"]:
        row = [news["what_happened"][:150]]
        for c in companies:
            row.append(next(f["score"] for f in feed if f["news_id"] == news["id"] and f["company_id"] == c["synthetic_u"]))
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EF3124")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[1:]:
            color = "C6EFCE" if cell.value >= FOCUS else "FFEB9C" if cell.value >= ALL else "F2F2F2"
            cell.fill = PatternFill("solid", fgColor=color)
    ws.column_dimensions["A"].width = 70
    for i in range(2, len(companies) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 24

    ws2 = wb.create_sheet("Обоснование")
    ws2.append(["Новость", "Компания", "Балл", "Лента", "Правила", "Почему важно", "Продукты банка"])
    for f in feed:
        ws2.append([f["news_id"], f["company_id"], f["score"], f["bucket"], ", ".join(f["rules"]),
                    "\n".join(f["why"]), ", ".join(f["products"])])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate([9, 12, 7, 14, 32, 70, 42], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(DATA / "feed_preview.xlsx")

    for news in golden["news"]:
        scores = [f["score"] for f in feed if f["news_id"] == news["id"]]
        print(news["id"], scores)


if __name__ == "__main__":
    main()
